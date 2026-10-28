#!/usr/bin/env python3
"""The question bank: one CSV per category, checked against the Question Bible and
turned into an Excel workbook the content team reviews in.

  python3 docs/content/question-bank/questionbank.py check   # rules and duplicates
  python3 docs/content/question-bank/questionbank.py build   # sync reviews, then rebuild the workbook

`build` first copies the team's review columns (review_status, reviewer, review_notes) from the
existing workbook back into the CSVs, so reviews made in Excel survive every rebuild.
"""
import csv
import difflib
import re
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
SEED = HERE.parent / "seed-questions.csv"
WORKBOOK = HERE / "question-bank.xlsx"

# File stem, sheet name, id prefix and display name for each category (doc 04 §1, plus sport).
CATEGORIES = [
    ("slang", "Slang", "SLG", "Slang & Pidgin"),
    ("nollywood", "Nollywood", "NOL", "Nollywood"),
    ("music", "Music", "MUS", "Music"),
    ("food", "Food", "FOD", "Food"),
    ("mixed", "Mixed", "MIX", "Mixed Naija"),
    ("sport", "Sport", "SPT", "Sport"),
]
COLUMNS = [
    "id", "category", "subcategory", "type", "difficulty", "question",
    "option_a", "option_b", "option_c", "option_d", "correct_option", "fact",
    "family_safe", "valid_until", "region_tag", "status", "source_url", "source_title",
    "added_on", "review_status", "reviewer", "review_notes",
]
REVIEW_COLUMNS = ["review_status", "reviewer", "review_notes"]
REVIEW_STATUSES = ["to check", "approved", "fix needed", "rejected"]
REGIONS = {"ALL", "SW", "SE", "SS", "N", "NC"}
TARGET_PER_CATEGORY = 600  # approved questions per category at launch (doc 04 §1)


def csv_path(stem):
    return HERE / f"{stem}.csv"


def read_rows(path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def write_rows(path, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in COLUMNS})


def normalise(text):
    text = text.lower().replace("as of october 2026,", "")
    return re.sub(r"[^a-z0-9 ]+", "", text).strip()


def check():
    errors, warnings = [], []
    seen_ids = {}
    all_questions = [(r["id"], normalise(r["question"])) for r in read_rows(SEED)]
    for stem, _sheet, prefix, _name in CATEGORIES:
        for n, r in enumerate(read_rows(csv_path(stem)), start=2):
            where = f"{stem}.csv line {n} ({r.get('id', '?')})"
            missing = [c for c in COLUMNS if c not in r]
            if missing:
                errors.append(f"{where}: missing columns {missing}")
                continue
            rid = r["id"]
            if not rid.startswith(prefix + "-"):
                errors.append(f"{where}: id should start with {prefix}-")
            if rid in seen_ids:
                errors.append(f"{where}: duplicate id, also in {seen_ids[rid]}")
            seen_ids[rid] = where
            if r["category"] != stem:
                errors.append(f"{where}: category should be {stem}")
            if r["difficulty"] not in {"1", "2", "3", "4", "5"}:
                errors.append(f"{where}: difficulty must be 1 to 5")
            q = r["question"]
            if len(q) > 120:
                errors.append(f"{where}: question is {len(q)} characters (max 120)")
            if not q.rstrip().endswith(("?", "…")):
                warnings.append(f"{where}: question should end with ? or …")
            options = [r[f"option_{c}"] for c in "abcd"]
            if r["type"] == "mcq4":
                if any(not o.strip() for o in options):
                    errors.append(f"{where}: a 4-option question needs all four options")
                if len({o.strip().lower() for o in options}) < 4:
                    errors.append(f"{where}: options repeat")
                if r["correct_option"] != "a":
                    errors.append(f"{where}: the correct answer goes in option_a (correct_option a)")
                for o in options:
                    if len(o) > 30:
                        errors.append(f"{where}: option '{o}' is {len(o)} characters (max 30)")
            elif r["type"] == "closest_number":
                if any(o.strip() for o in options):
                    errors.append(f"{where}: closest_number questions have no options")
                try:
                    float(r["correct_option"])
                except ValueError:
                    errors.append(f"{where}: closest_number needs a number in correct_option")
            else:
                errors.append(f"{where}: type must be mcq4 or closest_number")
            if len(r["fact"]) > 160:
                errors.append(f"{where}: fact is {len(r['fact'])} characters (max 160)")
            if not r["source_url"].startswith(("http://", "https://")):
                errors.append(f"{where}: needs a source_url")
            if r["valid_until"] and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", r["valid_until"]):
                errors.append(f"{where}: valid_until must look like 2027-10-31")
            if q.lower().startswith("as of") and not r["valid_until"]:
                errors.append(f"{where}: an 'As of' question needs valid_until")
            if r["family_safe"] not in {"true", "false"}:
                errors.append(f"{where}: family_safe must be true or false")
            if r["region_tag"] not in REGIONS:
                errors.append(f"{where}: region_tag must be one of {sorted(REGIONS)}")
            if r["review_status"] and r["review_status"] not in REVIEW_STATUSES:
                errors.append(f"{where}: review_status must be one of {REVIEW_STATUSES}")
            text = normalise(q)
            for other_id, other in all_questions:
                if text == other or difflib.SequenceMatcher(None, text, other).ratio() > 0.9:
                    warnings.append(f"{where}: looks like a duplicate of {other_id}")
                    break
            all_questions.append((rid, text))
    for w in warnings:
        print("warning:", w)
    for e in errors:
        print("error:", e)
    total = sum(len(read_rows(csv_path(s))) for s, *_ in CATEGORIES)
    print(f"{total} questions checked: {len(errors)} errors, {len(warnings)} warnings")
    return not errors


def sync_reviews_from_workbook():
    """Copies the review columns from the workbook into the CSVs (the workbook is where people review)."""
    if not WORKBOOK.exists():
        return 0
    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=True)
    reviews = {}
    for _stem, sheet, *_ in CATEGORIES:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        if "id" not in header:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(header, row))
            if rec.get("id"):
                reviews[rec["id"]] = {c: (rec.get(c) or "") for c in REVIEW_COLUMNS}
    changed = 0
    for stem, *_ in CATEGORIES:
        rows = read_rows(csv_path(stem))
        dirty = False
        for r in rows:
            rv = reviews.get(r["id"])
            if rv:
                for c in REVIEW_COLUMNS:
                    value = str(rv[c]).strip()
                    if value and value != r.get(c, ""):
                        r[c] = value
                        dirty = True
                        changed += 1
            if not r.get("review_status"):
                r["review_status"] = "to check"
                dirty = True
        if dirty:
            write_rows(csv_path(stem), rows)
    return changed


def build():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    synced = sync_reviews_from_workbook()
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="008751")
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    input_font = Font(name="Arial", size=10, color="0000FF")
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    widths = {"id": 11, "subcategory": 16, "type": 13, "difficulty": 9, "question": 48, "fact": 48,
              "source_url": 30, "source_title": 28, "review_notes": 30, "valid_until": 12}
    col = {name: get_column_letter(i + 1) for i, name in enumerate(COLUMNS)}

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    for _stem, sheet, *_ in CATEGORIES:
        ws = wb.create_sheet(sheet)
        rows = read_rows(csv_path(_stem))
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            values = []
            for c in COLUMNS:
                v = r.get(c, "")
                values.append(int(v) if c == "difficulty" and v.isdigit() else v)
            ws.append(values)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = base
                cell.alignment = Alignment(vertical="top", wrap_text=COLUMNS[cell.column - 1] in ("question", "fact", "review_notes"))
                name = COLUMNS[cell.column - 1]
                if name in REVIEW_COLUMNS:
                    cell.fill = review_fill
                if name == "source_url" and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = link_font
        for i, name in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 2)}"
        dv = DataValidation(type="list", formula1='"' + ",".join(REVIEW_STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col['review_status']}2:{col['review_status']}5000")

    # Summary: counts by review status, progress to the launch target, and the difficulty mix.
    summary["A1"] = "Who Sabi? question bank"
    summary["A1"].font = Font(name="Arial", size=14, bold=True)
    summary["A2"] = f"Built {date.today().isoformat()} from the CSV files in docs/content/question-bank/. Counts update when you change a review_status."
    summary["A2"].font = base
    headers = ["Category", "Questions", "To check", "Approved", "Fix needed", "Rejected", "Launch target", "Approved vs target"]
    summary.append([])
    summary.append(headers)
    hr = summary.max_row
    for cell in summary[hr]:
        cell.font, cell.fill = head_font, head_fill
    rs = col["review_status"]
    first = hr + 1
    for _stem, sheet, _prefix, name in CATEGORIES:
        r = summary.max_row + 1
        summary.append([
            name,
            f"=COUNTA({sheet}!A:A)-1",
            f'=COUNTIF({sheet}!{rs}:{rs},"to check")',
            f'=COUNTIF({sheet}!{rs}:{rs},"approved")',
            f'=COUNTIF({sheet}!{rs}:{rs},"fix needed")',
            f'=COUNTIF({sheet}!{rs}:{rs},"rejected")',
            TARGET_PER_CATEGORY,
            f"=IF(G{r}=0,0,D{r}/G{r})",
        ])
        summary[f"G{r}"].font = input_font
        summary[f"H{r}"].number_format = "0.0%"
    last = summary.max_row
    total = last + 1
    summary.append(["Total"] + [f"=SUM({c}{first}:{c}{last})" for c in "BCDEFG"] + [f"=IF(G{total}=0,0,D{total}/G{total})"])
    summary[f"H{total}"].number_format = "0.0%"
    for row in summary.iter_rows(min_row=first, max_row=total):
        for cell in row:
            if cell.column_letter == "G" and cell.row != total:
                cell.font = input_font
            else:
                cell.font = bold if cell.row == total else base
    summary[f"G{first}"].comment = Comment("Target from docs/04 §1: 600 approved questions per category at launch. Blue = an input you can change.", "Who Sabi?")

    summary.append([])
    summary.append(["Difficulty mix (all questions)"] + [name for *_x, name in CATEGORIES] + ["Doc 04 target"])
    dh = summary.max_row
    for cell in summary[dh]:
        cell.font, cell.fill = head_font, head_fill
    dcol = col["difficulty"]
    targets = {1: 0.15, 2: 0.20, 3: 0.30, 4: 0.20, 5: 0.15}
    for level in range(1, 6):
        summary.append([f"Level {level}"] + [f"=COUNTIF({sheet}!{dcol}:{dcol},{level})" for _s, sheet, *_ in CATEGORIES] + [targets[level]])
        r = summary.max_row
        summary[f"H{r}"].number_format = "0%"
        summary[f"H{r}"].font = input_font
        for cell in summary[r][:7]:
            cell.font = base

    summary.append([])
    legend = [
        "How to review",
        "1. Open a category sheet. Each row is one question; the right answer is always option_a (the game shuffles options).",
        "2. Open the source_url and check the fact, the spelling and that only one answer is right.",
        "3. Fill in only the yellow columns: review_status (approved, fix needed or rejected), reviewer, review_notes.",
        "   Example: approved | Ada | Checked on the CAF site, spelling fine.",
        "4. Upload the saved workbook to GitHub (same name, same folder) or send it back; the next build copies your marks into the CSVs.",
        "Only approved questions go into the game. Anything starting \"As of October 2026\" expires on its valid_until date.",
    ]
    for i, line in enumerate(legend):
        summary.append([line])
        summary[f"A{summary.max_row}"].font = bold if i == 0 else base
    summary.column_dimensions["A"].width = 30
    for c in "BCDEFGH":
        summary.column_dimensions[c].width = 15
    wb.save(WORKBOOK)
    print(f"wrote {WORKBOOK.name}; {synced} review values copied from the previous workbook")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "check"
    if cmd == "check":
        sys.exit(0 if check() else 1)
    if cmd == "build":
        if not check():
            sys.exit("fix the errors above before building")
        build()
        sys.exit(0)
    sys.exit(f"unknown command {cmd!r}: use check or build")
